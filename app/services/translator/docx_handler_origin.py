import os
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from typing import Optional, List, Tuple, Dict
from .core import HYMTTranslator
import zipfile
from pathlib import Path


class DocxHandler:
    """Word 文档 (.docx) 处理器（支持：原顺序+合并单元格+嵌套表格）"""

    def __init__(self, translator: HYMTTranslator):
        self.translator = translator

    def _process_single_table_to_html(self, table, target_lang, show_bilingual) -> str:
        """
        处理单个表格转HTML（支持合并单元格+嵌套表格）
        :param table: docx Table 对象
        :param target_lang: 目标语言
        :param show_bilingual: 是否显示中英对照
        :return: 带合并单元格+嵌套表格的表格HTML
        """
        # 解析表格的合并单元格信息
        merge_info = self._parse_table_merge_info(table)
        html_rows = []

        # 逐行处理表格
        for row_idx, row in enumerate(table.rows):
            html_cells = []
            for col_idx, cell in enumerate(row.cells):
                # 跳过被合并的单元格（已被其他单元格覆盖）
                if (row_idx, col_idx) in merge_info['merged_cells']:
                    continue

                # 获取单元格合并属性
                colspan = merge_info['colspan'].get((row_idx, col_idx), 1)
                rowspan = merge_info['rowspan'].get((row_idx, col_idx), 1)

                # 核心修改：解析单元格内的内容（文本 + 嵌套表格）
                cell_content_html = self._parse_cell_content(cell, target_lang, show_bilingual)

                # 生成带合并属性的单元格HTML
                cell_attrs = []
                if colspan > 1:
                    cell_attrs.append(f'colspan="{colspan}"')
                if rowspan > 1:
                    cell_attrs.append(f'rowspan="{rowspan}"')
                attrs_str = ' '.join(cell_attrs) if cell_attrs else ''

                cell_html = f'<td {attrs_str}>{cell_content_html}</td>' if attrs_str else f'<td>{cell_content_html}</td>'
                html_cells.append(cell_html)

            # 生成行HTML
            html_rows.append(f'<tr>{"".join(html_cells)}</tr>')

        # 生成完整表格HTML（优化样式适配嵌套表格）
        table_html = f"""
        <div class="document-table" style="margin: 20px 0; overflow-x: auto;">
            <table border="1" cellpadding="8" cellspacing="0" style="border-collapse: collapse; width: 100%;">
                {"".join(html_rows)}
            </table>
        </div>
        """
        return table_html

    def _parse_cell_content(self, cell, target_lang, show_bilingual) -> str:
        """
        解析单元格内容（文本 + 嵌套表格）
        :param cell: docx Cell 对象
        :param target_lang: 目标语言
        :param show_bilingual: 是否显示中英对照
        :return: 单元格内容的 HTML 字符串
        """
        content_parts = []

        # 1. 提取单元格内的文本
        cell_text = cell.text.strip()
        if cell_text:
            # 检测是否为中文，如果是则跳过翻译
            if self._is_chinese_text(cell_text):
                # 中文不翻译，原文和译文相同
                text_html = f"""
                <div class="cell-text">
                    <div class="original-text">{cell_text}</div>
                    <div class="translated-text">{cell_text}</div>
                </div>
                """
                content_parts.append(text_html)
            elif show_bilingual:
                translated_text = self.translator.translate_text(cell_text, target_lang)
                text_html = f"""
                <div class="cell-text">
                    <div class="original-text">{cell_text}</div>
                    <div class="translated-text">{translated_text}</div>
                </div>
                """
                content_parts.append(text_html)
            else:
                translated_text = self.translator.translate_text(cell_text, target_lang)
                text_html = f'<div class="cell-text translated-text">{translated_text}</div>'
                content_parts.append(text_html)

        # 2. 检查并解析单元格内的嵌套表格（核心新增逻辑）
        # 遍历单元格的 XML 节点，查找嵌套的<table>节点
        cell_elem = cell._tc
        nested_tables = cell_elem.xpath('.//w:tbl')
        if nested_tables:
            # 遍历所有嵌套表格，递归处理
            for nested_tbl_elem in nested_tables:
                # 从 XML 节点构建 docx Table 对象
                from docx.table import Table
                nested_table = Table(nested_tbl_elem, cell)
                # 递归生成嵌套表格的 HTML
                nested_table_html = self._process_single_table_to_html(nested_table, target_lang, show_bilingual)
                # 添加嵌套表格样式（缩进 + 边框区分）
                styled_nested_table = f"""
                <div class="nested-table" style="margin: 10px 0 10px 20px; border: 1px dashed #666; padding: 5px;">
                    {nested_table_html}
                </div>
                """
                content_parts.append(styled_nested_table)

        # 3. 处理空单元格
        if not content_parts:
            return '<div class="empty-cell">&nbsp;</div>'

        return ''.join(content_parts)


    def _parse_table_merge_info(self, table) -> Dict[str, any]:
        """
        解析表格的合并单元格信息（保留原有逻辑）
        """
        colspan = {}
        rowspan = {}
        merged_cells = set()

        # 遍历所有单元格，解析合并属性
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                # 跳过已标记为合并的单元格
                if (row_idx, col_idx) in merged_cells:
                    continue

                # 解析横向合并（gridSpan）
                cell_elem = cell._tc
                grid_span = cell_elem.xpath('.//w:gridSpan')
                if grid_span:
                    span = int(grid_span[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'))
                    colspan[(row_idx, col_idx)] = span
                    # 标记被合并的列
                    for c in range(1, span):
                        merged_cells.add((row_idx, col_idx + c))

                # 解析纵向合并（vMerge）
                v_merge = cell_elem.xpath('.//w:vMerge')
                if v_merge:
                    v_merge_val = v_merge[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                    # vMerge="restart" 表示合并起始单元格，需要计算合并行数
                    if v_merge_val == 'restart':
                        merge_rows = 1
                        # 向下遍历，统计连续的合并行
                        for r in range(row_idx + 1, len(table.rows)):
                            next_cell = table.rows[r].cells[col_idx]
                            next_v_merge = next_cell._tc.xpath('.//w:vMerge')
                            if next_v_merge and next_v_merge[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') is None:
                                merge_rows += 1
                                merged_cells.add((r, col_idx))
                            else:
                                break
                        rowspan[(row_idx, col_idx)] = merge_rows

        return {
            'colspan': colspan,
            'rowspan': rowspan,
            'merged_cells': merged_cells
        }

    # 保留原有 convert_to_html/_parse_document_elements/_process_single_paragraph_to_html 等核心方法
    def convert_to_html(
            self,
            input_path: str,
            output_dir: str,
            target_lang: str = "Chinese",
            show_bilingual: bool = True,
            translate_all: int = 0,
            preserve_original_styles: bool = True
    ) -> str:
        """
        将 Word 文档转换为 HTML 并翻译（支持：原顺序 + 合并单元格 + 嵌套表格）
        """
        os.makedirs(output_dir, exist_ok=True)
        print(f"Processing Word to HTML (full flow): {input_path}")
        doc = Document(input_path)

        # 构建 HTML 内容（增强样式适配）
        html_content = [self._get_html_header(preserve_original_styles)]

        # 进度追踪
        tracker = self.translator.get_progress_tracker()
        # 解析文档所有原生元素（段落/表格/图片）并计数
        all_elements = self._parse_document_elements(doc, input_path)
        total_elements = len(all_elements)
        elements_to_process = total_elements if translate_all == 0 else min(translate_all, total_elements)
        tracker.set_file_info(os.path.basename(input_path), elements_to_process, "element")

        # 核心：按原生顺序处理每个元素
        processed_count = 0
        img_idx = 0  # 图片索引（用于命名）

        # 修改：为每个文档创建独立的图片文件夹
        base_name = os.path.basename(input_path)
        name_without_ext = os.path.splitext(base_name)[0]
        img_folder_name = f"{name_without_ext}_images"
        img_dir = os.path.join(output_dir, img_folder_name)
        os.makedirs(img_dir, exist_ok=True)

        for elem_type, elem_data in all_elements:
            if translate_all > 0 and processed_count >= elements_to_process:
                break

            if elem_type == "paragraph":
                para = elem_data
                html = self._process_single_paragraph_to_html(para, target_lang, show_bilingual,
                                                              preserve_original_styles)
                if html:
                    html_content.append(html)
                processed_count += 1
                tracker.update_paragraph(processed_count)

            elif elem_type == "table":
                table = elem_data
                html = self._process_single_table_to_html(table, target_lang, show_bilingual)
                html_content.append(html)
                processed_count += 1
                tracker.update_paragraph(processed_count)

            elif elem_type == "image":
                # 处理单张图片（按出现顺序）
                img_file = elem_data
                img_ext = Path(img_file).suffix
                img_name = f"image_{img_idx}{img_ext}"
                img_save_path = os.path.join(img_dir, img_name)
                # 提取图片文件
                with zipfile.ZipFile(input_path) as doc_zip:
                    with open(img_save_path, 'wb') as f:
                        f.write(doc_zip.read(img_file))
                # 生成图片 HTML（使用独立文件夹路径）
                img_rel_path = os.path.join(img_folder_name, img_name)
                img_html = f'''
                <div class="document-image">
                    <img src="{self._escape_html(img_rel_path)}" alt="文档图片 {img_idx + 1}" 
                         style="max-width: 100%; height: auto; margin: 20px 0;">
                </div>
                '''
                html_content.append(img_html)
                img_idx += 1
                processed_count += 1
                tracker.update_paragraph(processed_count)

        # 闭合 HTML 标签
        html_content.append(self._get_html_footer())

        # 保存 HTML 文件
        output_path = os.path.join(output_dir, f"{name_without_ext}_translated.html")

        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(html_content))

        tracker.mark_completed()
        print(f"翻译后的 HTML（完整样式 + 原顺序 + 合并单元格 + 嵌套表格）已保存至：{output_path}")
        return output_path

    def _parse_document_elements(self, doc: Document, doc_path: str) -> List[Tuple[str, object]]:
        """解析DOCX文档的所有元素，按原生顺序返回"""
        elements = []
        # 1. 先解析段落和表格（基于docx库的节点遍历）
        for child in doc.element.body:
            tag = child.tag.rpartition('}')[-1]  # 获取节点标签（p/tbl）
            if tag == 'p':
                # 找到对应的Paragraph对象
                para_idx = len([e for e in elements if e[0] == "paragraph"])
                elements.append(("paragraph", doc.paragraphs[para_idx]))
            elif tag == 'tbl':
                # 找到对应的Table对象
                table_idx = len([e for e in elements if e[0] == "table"])
                elements.append(("table", doc.tables[table_idx]))

        # 2. 解析图片（关联到段落位置）
        # 先提取所有图片文件
        doc_zip = zipfile.ZipFile(doc_path)
        image_files = [f for f in doc_zip.namelist() if f.startswith('word/media/')]
        doc_zip.close()

        # 遍历段落，找到包含图片的段落，插入到对应位置
        img_idx = 0
        new_elements = []
        for elem in elements:
            new_elements.append(elem)
            # 检查当前段落是否包含图片
            if elem[0] == "paragraph":
                para = elem[1]
                if para._element.xpath('.//w:drawing') and img_idx < len(image_files):
                    new_elements.append(("image", image_files[img_idx]))
                    img_idx += 1

        return new_elements

    def _process_single_paragraph_to_html(self, para, target_lang, show_bilingual, preserve_styles):
        """处理单个段落转 HTML（复用原有逻辑）"""
        text = para.text.strip()
        if not text:
            return '<br/>'

        # 识别段落样式
        para_style = self._get_paragraph_style_type(para)

        # 检测是否为中文，如果是则跳过翻译
        if self._is_chinese_text(text):
            # 中文段落不翻译，直接返回原文
            return self._generate_paragraph_html(text, text, para_style, preserve_styles)

        if show_bilingual:
            translated = self.translator.translate_text(text, target_lang)
            return self._generate_paragraph_html(text, translated, para_style, preserve_styles)
        else:
            translated = self.translator.translate_text(text, target_lang)
            return self._generate_paragraph_html("", translated, para_style, preserve_styles)

    def _is_chinese_text(self, text: str) -> bool:
        """
        检测文本是否主要为中文（包括阿拉伯数字和常见标点）
        :param text: 待检测文本
        :return: True 表示主要是中文、数字或标点，False 表示非中文
        """
        import re
        # 统计中文字符、阿拉伯数字和中文标点的数量
        chinese_chars = re.findall(r'[\u4e00-\u9fff]', text)
        arabic_numerals = re.findall(r'[0-9]', text)
        # 中文标点符号范围：包括常用标点如，。！？；：""''（）【】《》等
        chinese_punctuation = re.findall(r'[\u3000-\u303f\uff00-\uffef]', text)

        # 如果中文字符、阿拉伯数字或中文标点占比超过 70%，则认为是中文文本
        if len(text) > 0:
            chinese_or_digit_count = len(chinese_chars) + len(arabic_numerals) + len(chinese_punctuation)
            ratio = chinese_or_digit_count / len(text)
            return ratio >= 0.7
        return False

    # 辅助方法（优化样式适配嵌套表格）
    def _get_html_header(self, preserve_styles: bool) -> str:
        """生成HTML头部（补充嵌套表格样式）"""
        style = """
        <style>
            body { font-family: Arial, sans-serif; line-height: 1.6; padding: 20px; }
            .document-table table { border: 1px solid #000; }
            .document-table td { border: 1px solid #000; vertical-align: top; padding: 8px; }
            /* 嵌套表格样式 */
            .nested-table { border: 1px dashed #666 !important; margin: 10px; padding: 5px; }
            .nested-table .document-table { margin: 0 !important; }
            /* 单元格文本样式 */
            .cell-text { margin-bottom: 5px; }
            .original-text { color: #333; margin-bottom: 3px; }
            .translated-text { color: #0066cc; }
            .empty-cell { color: #ccc; }
            .document-image { text-align: center; }
            h1, h2, h3 { margin: 15px 0; }
            ul, ol { margin: 10px 0 10px 20px; }
        </style>
        """ if preserve_styles else ""
        return f"""
        <!DOCTYPE html>
        <html lang="zh-CN">
        <head>
            <meta charset="UTF-8">
            <title>Translated Document</title>
            {style}
        </head>
        <body>
        """

    def _get_html_footer(self) -> str:
        """生成HTML尾部"""
        return """
        </body>
        </html>
        """

    def _escape_html(self, text: str) -> str:
        """转义HTML特殊字符"""
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")

    def _get_paragraph_style_type(self, para) -> str:
        """识别段落样式类型（标题/列表/普通文本）"""
        style_name = para.style.name.lower()
        if 'heading 1' in style_name or '标题1' in style_name:
            return 'h1'
        elif 'heading 2' in style_name or '标题2' in style_name:
            return 'h2'
        elif 'heading 3' in style_name or '标题3' in style_name:
            return 'h3'
        elif para.style.name in ['List Paragraph', '列表段落']:
            return 'list'
        else:
            return 'p'

    def _generate_paragraph_html(self, original: str, translated: str, style_type: str, preserve_styles: bool) -> str:
        """生成段落HTML（保留样式）"""
        if not preserve_styles:
            style_type = 'p'

        content = ""
        if original:
            content += f'<span class="original-text">{original}</span><br/>'
        content += f'<span class="translated-text">{translated}</span>'

        if style_type == 'h1':
            return f'<h1>{content}</h1>'
        elif style_type == 'h2':
            return f'<h2>{content}</h2>'
        elif style_type == 'h3':
            return f'<h3>{content}</h3>'
        elif style_type == 'list':
            return f'<li>{content}</li>'
        else:
            return f'<p>{content}</p>'